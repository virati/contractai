"""
Create a sample contract document for testing purposes.
"""

from docx import Document
import os


def create_sample_contract():
    """Create a sample contract document."""
    doc = Document()

    # Add title
    doc.add_heading("SERVICE AGREEMENT", 0)

    # Add contract content
    doc.add_heading("1. PARTIES", level=1)
    doc.add_paragraph(
        "This Service Agreement ('Agreement') is entered into as of January 1, 2024, "
        "between TechCorp Inc. ('Company') and John Doe ('Contractor')."
    )

    doc.add_heading("2. SCOPE OF WORK", level=1)
    doc.add_paragraph(
        "The Contractor agrees to provide software development services as specified "
        "in the Statement of Work attached as Exhibit A. The Contractor shall deliver "
        "all work products by the agreed-upon deadlines and meet the quality standards "
        "specified by the Company."
    )

    doc.add_heading("3. PAYMENT TERMS", level=1)
    doc.add_paragraph(
        "The Company agrees to pay the Contractor $100 per hour for services rendered. "
        "Payment shall be made within 30 days of receipt of invoice. The Contractor "
        "shall submit invoices monthly with detailed time sheets."
    )

    doc.add_heading("4. INTELLECTUAL PROPERTY", level=1)
    doc.add_paragraph(
        "All work products, deliverables, and intellectual property created by the "
        "Contractor in the course of performing services under this Agreement shall "
        "be the sole property of the Company. The Contractor hereby assigns all rights, "
        "title, and interest in such work products to the Company."
    )

    doc.add_heading("5. CONFIDENTIALITY", level=1)
    doc.add_paragraph(
        "The Contractor agrees to maintain the confidentiality of all Company "
        "proprietary information and trade secrets. This obligation shall survive "
        "termination of this Agreement for a period of five (5) years."
    )

    doc.add_heading("6. LIABILITY AND INDEMNIFICATION", level=1)
    doc.add_paragraph(
        "The Contractor shall indemnify and hold harmless the Company from any claims, "
        "damages, or expenses arising from the Contractor's performance of services "
        "under this Agreement. The Contractor's liability under this Agreement shall "
        "not exceed the total amount paid to the Contractor in the preceding six (6) months."
    )

    doc.add_heading("7. TERMINATION", level=1)
    doc.add_paragraph(
        "Either party may terminate this Agreement with thirty (30) days written notice. "
        "Upon termination, the Contractor shall deliver all work products and return "
        "all Company property. The Company shall pay for all services performed up to "
        "the termination date."
    )

    doc.add_heading("8. INDEPENDENT CONTRACTOR STATUS", level=1)
    doc.add_paragraph(
        "The Contractor is an independent contractor and not an employee of the Company. "
        "The Contractor is responsible for all taxes and shall not be entitled to "
        "employee benefits."
    )

    # Save the document
    output_path = "examples/sample_contract.docx"
    doc.save(output_path)
    print(f"Sample contract created at: {output_path}")
    return output_path


if __name__ == "__main__":
    create_sample_contract()
